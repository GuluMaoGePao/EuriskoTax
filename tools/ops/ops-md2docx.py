# -*- coding: utf-8 -*-
"""Markdown -> DOCX 导出工具（docs/marketing 对外发送版）。

用法（项目根目录执行）：

    python tools\\ops\\ops-md2docx.py docs\\marketing\\business-plan-for-partners.md
    python tools\\ops\\ops-md2docx.py <input.md> <output.docx>

约定：

- **真源永远是 .md**，.docx 是导出件；改内容改 .md 后重新导出，不要在 Word 里改。
- 封面副标题（可选）：md 中任意位置写 `<!-- subtitle: 一行副标题 -->`。
- 封面底部提示（可选）：`<!-- notice: 一句话 -->`，默认「内部资料 · 请勿外传」。
- 导出时剔除内部元信息：用 `<!-- export:skip -->` ... `<!-- /export:skip -->` 包住
  （例如"文档分工"表这类只给内部看的内容，md 里保留、Word 里不出现）。

依赖：`pip install python-docx`（导出用，主项目运行不需要）。
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

LATIN = "Segoe UI"
CJK = "微软雅黑"
MONO = "Consolas"
CONTENT_W = 16.0  # cm（A4 21 - 2.5 * 2 页边距）

# emoji 在部分机器上会缺字，统一替换为可用符号
EMOJI_MAP = [
    ("\u2705", "\u221a"),      # ✅ -> √
    ("\u274c", "\u00d7"),      # ❌ -> ×
    ("\u23f3", "\u25cb"),      # ⏳ -> ○
    ("\u2b50", "\u2605"),      # ⭐ -> ★
    ("\U0001f6a9", "\u203b"),  # 🚩 -> ※
    ("\U0001f6a8", "\u203b"),  # 🚨 -> ※
    ("\u26a0", "\u203b"),      # ⚠  -> ※
    ("\ufe0f", ""),            # 变体选择符
]

SKIP_BLOCK = re.compile(
    r"<!--\s*export:skip\s*-->.*?<!--\s*/export:skip\s*-->", re.S)
SUBTITLE = re.compile(r"<!--\s*subtitle:\s*(.+?)\s*-->")
NOTICE = re.compile(r"<!--\s*notice:\s*(.+?)\s*-->")
TOK = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]*\]\([^)]*\)|\*[^*\s][^*]*\*)")
SEP_ROW = re.compile(r"^\|[\s:\-|]+\|$")
DEFAULT_NOTICE = "内部资料 · 请勿外传"


def clean(s):
    for a, b in EMOJI_MAP:
        s = s.replace(a, b)
    return s


def plain(s):
    """去掉行内 Markdown，用于目录条目等纯文本位置。"""
    s = re.sub(r"\[([^\]]*)\]\([^)]*\)", r"\1", s)
    s = s.replace("**", "").replace("`", "")
    s = re.sub(r"\*([^*]+)\*", r"\1", s)
    return clean(s)


def eff_len(s):
    """估算显示宽度（CJK 记 2），用于表格列宽自适应。"""
    s = re.sub(r"\*\*|`|\[|\]|\([^)]*\)", "", s)
    return sum(2 if ord(ch) > 0x2E7F else 1 for ch in s)


# ------------------------------------------------------------------ 文本渲染
def _fonts(run, latin, ea):
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), ea)


def sr(run, size=10.5, bold=False, italic=False, color=None, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    _fonts(run, MONO, MONO) if mono else _fonts(run, LATIN, CJK)
    return run


def add_link(p, url, text, size=9):
    r_id = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hl = OxmlElement("w:hyperlink")
    hl.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    f = OxmlElement("w:rFonts")
    f.set(qn("w:ascii"), LATIN)
    f.set(qn("w:hAnsi"), LATIN)
    f.set(qn("w:eastAsia"), CJK)
    rPr.append(f)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rPr.append(c)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    r.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = clean(text)
    r.append(t)
    hl.append(r)
    p._p.append(hl)


NESTED = re.compile(r"\*\*|`|\[[^\]]*\]\([^)]*\)")


def add_inline(p, text, size=10.5, bold=False, color=None, mono=False,
               italic=False):
    for part in TOK.split(clean(text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            inner = part[2:-2]
            # 加粗内部可能还嵌着行内代码/链接（如 **同时注册 `.cn` 并 301**）
            if NESTED.search(inner):
                add_inline(p, inner, size, True, color, mono)
            else:
                sr(p.add_run(inner), size, True, italic, color, mono)
        elif part.startswith("`") and part.endswith("`") and len(part) >= 2:
            sr(p.add_run(part[1:-1]), size - 0.5, False, False, "C0392B", True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            inner = part[1:-1]
            if NESTED.search(inner):
                add_inline(p, inner, size, bold, color, mono, True)
            else:
                sr(p.add_run(inner), size, bold, True, color, mono)
        elif part.startswith("["):
            m = re.match(r"\[([^\]]*)\]\(([^)]*)\)", part)
            label, url = m.group(1), m.group(2)
            if url.startswith("http"):
                add_link(p, url, label, size - 1)
            else:
                sr(p.add_run(label), size, bold, False, color, mono)
        else:
            sr(p.add_run(part), size, bold, italic, color, mono)


# ------------------------------------------------------------------ 段落排版
def set_para(p, before=0, after=5, line=1.4, left=0, right=0, hanging=None,
             align=None, keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    if left:
        pf.left_indent = Cm(left)
    if right:
        pf.right_indent = Cm(right)
    if hanging is not None:
        pf.first_line_indent = Cm(-hanging)
    if align is not None:
        pf.alignment = align
    if keep_next:
        pf.keep_with_next = True
    return p


def _ppr(p, elm, successors):
    p._p.get_or_add_pPr().insert_element_before(elm, *successors)


def shade_para(p, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    _ppr(p, shd, ("w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr", "w:sectPr"))


def side_border(p, tag, color, size):
    pBdr = OxmlElement("w:pBdr")
    e = OxmlElement("w:" + tag)
    e.set(qn("w:val"), "single")
    e.set(qn("w:sz"), size)
    e.set(qn("w:space"), "8" if tag == "left" else "2")
    e.set(qn("w:color"), color)
    pBdr.append(e)
    _ppr(p, pBdr, ("w:shd", "w:tabs", "w:spacing", "w:ind", "w:jc", "w:rPr",
                   "w:sectPr"))


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_field(p, instr, size=8.5, color="808080"):
    r1 = p.add_run()
    sr(r1, size=size, color=color)
    fc = OxmlElement("w:fldChar")
    fc.set(qn("w:fldCharType"), "begin")
    r1._r.append(fc)

    r2 = p.add_run()
    sr(r2, size=size, color=color)
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = " %s " % instr
    r2._r.append(it)

    r3 = p.add_run()
    sr(r3, size=size, color=color)
    fs = OxmlElement("w:fldChar")
    fs.set(qn("w:fldCharType"), "separate")
    r3._r.append(fs)

    sr(p.add_run("1"), size=size, color=color)

    r5 = p.add_run()
    sr(r5, size=size, color=color)
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    r5._r.append(fe)


# ------------------------------------------------------------------ 文档骨架
def setup_styles(doc):
    st = doc.styles["Normal"]
    st.font.size = Pt(10.5)
    st.font.name = LATIN
    st.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK)

    for name, size, col in [("Heading 1", 14.5, "1F3864"),
                            ("Heading 2", 12, "2E5C8A"),
                            ("Heading 3", 11, "2E5C8A")]:
        s = doc.styles[name]
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(col)
        s.font.name = LATIN
        s.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), CJK)


def setup_section(doc, header_text):
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.5)
    sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.4)
    sec.bottom_margin = Cm(2.2)
    sec.different_first_page_header_footer = True

    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    sr(hp.add_run(clean(header_text)), size=8, color="A6A6A6")

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr(fp.add_run("第 "), size=8.5, color="808080")
    add_field(fp, "PAGE")
    sr(fp.add_run(" 页 / 共 "), size=8.5, color="808080")
    add_field(fp, "NUMPAGES")
    sr(fp.add_run(" 页"), size=8.5, color="808080")
    # 首页页脚留白（封面底部已有密级提示，避免重复）


def add_cover(doc, title, subtitle, meta_lines, notice):
    p = doc.add_paragraph()
    set_para(p, before=110, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    sr(p.add_run(clean(title)), size=26, bold=True, color="1F3864")

    if subtitle:
        p = doc.add_paragraph()
        set_para(p, before=0, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
        sr(p.add_run(clean(subtitle)), size=12, color="595959")

    p = doc.add_paragraph()
    set_para(p, before=0, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    side_border(p, "bottom", "2E75B6", "8")

    if meta_lines:
        render(doc, meta_lines, center=True)
        p = doc.add_paragraph()
        set_para(p, before=0, after=10)

    p = doc.add_paragraph()
    set_para(p, before=60, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    sr(p.add_run(clean(notice)), size=9, color="A6A6A6")

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_toc(doc, headings):
    p = doc.add_paragraph()
    set_para(p, before=0, after=12)
    side_border(p, "bottom", "BFBFBF", "6")
    sr(p.add_run("目  录"), size=17, bold=True, color="1F3864")

    if not headings:
        return
    first, last = True, None
    for lvl, text in headings:
        para = doc.add_paragraph()
        set_para(para, before=0, after=0, line=1.2, left=0 if lvl == 1 else 0.75)
        size = 10.5 if lvl == 1 else 9.5
        color = "1F3864" if lvl == 1 else "595959"
        if first:
            r = para.add_run()
            sr(r, size=size, bold=True, color=color)
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), "begin")
            r._r.append(fc)
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = ' TOC \\o "1-2" \\h \\z \\u '
            r._r.append(it)
            fs = OxmlElement("w:fldChar")
            fs.set(qn("w:fldCharType"), "separate")
            r._r.append(fs)
            first = False
        sr(para.add_run(plain(text)), size=size, bold=(lvl == 1), color=color)
        last = para

    r = last.add_run()
    sr(r, size=6)
    fe = OxmlElement("w:fldChar")
    fe.set(qn("w:fldCharType"), "end")
    r._r.append(fe)

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_table(doc, rows):
    ncols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False  # 写入 <w:tblLayout w:type="fixed"/>，勿再手动追加

    weights = []
    for ci in range(ncols):
        vals = sorted(eff_len(r[ci]) for r in rows if ci < len(r) and r[ci])
        if not vals:
            vals = [6]
        pick = vals[min(len(vals) - 1, int(len(vals) * 0.75))]
        weights.append(max(6.0, float(pick)))
    widths = [max(1.75, CONTENT_W * w / sum(weights)) for w in weights]
    widths = [w * CONTENT_W / sum(widths) for w in widths]

    fs = 8.5 if ncols <= 4 else 8.0
    for ri, r in enumerate(rows):
        row = t.add_row()
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        if ri == 0:
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            trPr.append(th)
        for ci in range(ncols):
            cell = row.cells[ci]
            cell.width = Cm(widths[ci])
            para = cell.paragraphs[0]
            set_para(para, before=1, after=1, line=1.2,
                     align=WD_ALIGN_PARAGRAPH.CENTER if ri == 0 else None)
            add_inline(para, r[ci] if ci < len(r) else "", size=fs,
                       bold=(ri == 0))
            if ri == 0:
                shade_cell(cell, "DCE6F1")

    sp = doc.add_paragraph()
    set_para(sp, before=0, after=0, line=1.0)
    sr(sp.add_run(""), size=4)


# ------------------------------------------------------------------ Markdown
def split_head(lines):
    """拆出 H1 标题与紧随其后的引用块（封面元信息），返回正文剩余行。"""
    title, meta, i = "", [], 0
    while i < len(lines):
        if lines[i].startswith("# "):
            title = lines[i][2:].strip()
            i += 1
            break
        i += 1
    while i < len(lines):
        s = lines[i].strip()
        if s == "":
            i += 1
            continue
        if s.startswith(">"):
            meta.append(s.lstrip(">").strip())
            i += 1
            continue
        break
    return title, meta, lines[i:]


def scan_headings(lines):
    out = []
    for ln in lines:
        s = ln.strip()
        if re.match(r"^###\s+", s):
            out.append((2, re.sub(r"^###\s+", "", s)))
        elif re.match(r"^##\s+", s):
            out.append((1, re.sub(r"^##\s+", "", s)))
    return out


def render(doc, lines, center=False):
    align = WD_ALIGN_PARAGRAPH.CENTER if center else None
    i, num, in_num = 0, 0, False
    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()
        if s == "":
            i += 1
            in_num = False
            continue

        if s.startswith("```"):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                buf.append(lines[i].rstrip())
                i += 1
            i += 1
            for k, ln in enumerate(buf):
                p = doc.add_paragraph()
                set_para(p, before=(4 if k == 0 else 0),
                         after=(5 if k == len(buf) - 1 else 0), line=1.15,
                         left=0 if center else 0.3, align=align)
                shade_para(p, "F2F2F2")
                txt = clean(ln)
                sr(p.add_run(txt if txt.strip() else " "), size=9, mono=True,
                   color="333333")
            in_num = False
            continue

        m = re.match(r"^(#{1,6})\s+(.*)$", s)
        if m:
            # md "##" -> Heading 1，"###" -> Heading 2：与目录域 \o "1-2" 一致
            lvl = min(3, max(1, len(m.group(1)) - 1))
            p = doc.add_paragraph(style="Heading %d" % lvl)
            set_para(p, before=(14 if lvl == 1 else 9),
                     after=(6 if lvl == 1 else 4), line=1.3, keep_next=True)
            add_inline(p, m.group(2), size={1: 14.5, 2: 12, 3: 11}[lvl],
                       bold=True, color="1F3864" if lvl == 1 else "2E5C8A")
            in_num = False
            i += 1
            continue

        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", s):
            i += 1
            continue

        if s.startswith(">"):
            buf = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                buf.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            set_para(p, before=5, after=5, line=1.35, left=0.4)
            side_border(p, "left", "2E75B6", "18")
            shade_para(p, "F4F8FD")
            for k, ln in enumerate(buf):
                if k:
                    r = p.add_run()
                    sr(r, size=10)
                    r.add_break()
                add_inline(p, ln, size=10, color="1F3864")
            in_num = False
            continue

        if s.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                ln = lines[i].strip()
                if not SEP_ROW.match(ln):
                    rows.append([c.strip() for c in ln.strip("|").split("|")])
                i += 1
            if rows:
                add_table(doc, rows)
            in_num = False
            continue

        m = re.match(r"^(\s*)[-*+]\s+(.*)$", raw)
        if m:
            level = min(2, len(m.group(1)) // 2)
            p = doc.add_paragraph(style="List Bullet")
            set_para(p, before=0, after=3, line=1.35,
                     left=0.75 + 0.65 * level, hanging=0.4)
            add_inline(p, m.group(2), size=10.5)
            in_num = False
            i += 1
            continue

        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", raw)
        if m:
            level = min(2, len(m.group(1)) // 2)
            num = num + 1 if in_num else 1
            in_num = True
            p = doc.add_paragraph()
            set_para(p, before=0, after=3, line=1.35,
                     left=0.85 + 0.65 * level, hanging=0.55)
            sr(p.add_run("%d. " % num), size=10.5, bold=True)
            add_inline(p, m.group(3), size=10.5)
            i += 1
            continue

        p = doc.add_paragraph()
        set_para(p, before=0, after=5, line=1.45, align=align)
        add_inline(p, s, size=10.5)
        in_num = False
        i += 1


def set_update_fields(doc):
    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))
    if zoom is not None and not zoom.get(qn("w:percent")):
        zoom.set(qn("w:percent"), "100")  # 默认模板缺这个必填属性
    upd = OxmlElement("w:updateFields")
    upd.set(qn("w:val"), "true")
    settings.insert_element_before(
        upd, "w:hdrShapeDefaults", "w:footnotePr", "w:endnotePr", "w:compat",
        "w:docVars", "w:rsids", "m:mathPr", "w:attachedSchema", "w:themeFontLang",
        "w:clrSchemeMapping", "w:decimalSymbol", "w:listSeparator")


def main(argv):
    if len(argv) < 2:
        print(__doc__)
        return 1
    src = argv[1]
    out = argv[2] if len(argv) > 2 else re.sub(r"\.md$", ".docx", src)

    with open(src, encoding="utf-8") as fh:
        text = fh.read()

    m = SUBTITLE.search(text)
    subtitle = m.group(1) if m else ""
    m = NOTICE.search(text)
    notice = m.group(1) if m else DEFAULT_NOTICE
    text = SKIP_BLOCK.sub("", text)
    text = SUBTITLE.sub("", text)
    text = NOTICE.sub("", text)

    lines = text.split("\n")
    title, meta, body = split_head(lines)

    doc = Document()
    setup_styles(doc)
    setup_section(doc, title)
    add_cover(doc, title, subtitle, meta, notice)
    add_toc(doc, scan_headings(body))
    render(doc, body)
    set_update_fields(doc)
    doc.save(out)

    print("OK  ->", out)
    print("title:", title)
    print("toc  :", len(scan_headings(body)), "entries /", len(meta), "meta lines")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
